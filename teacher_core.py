# -*- coding: utf-8 -*-
"""المنطق المشترك بين نسخة ويندوز (main.py) ونسخة الأندرويد (android/src/main.py).
ما فيه حتى استعمال لواجهة رسومية: ملفات، بحث، تعمير قوالب Word، واستخراج نص PDF."""

import copy
import json
import os
import re
import unicodedata

try:  # مكتبة تعمير القوالب (python-docx)
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:  # pragma: no cover
    docx = None

try:  # استخراج نص PDF (مولّد المذكرات)
    import pymupdf as fitz
except ImportError:  # pragma: no cover
    try:
        import fitz
    except ImportError:
        fitz = None

TEMPLATE_NAME = "template.docx"
INVALID_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
RESERVED_NAMES = {
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}
IGNORED_FILES = {"thumbs.db", "desktop.ini"}
_ARABIC_MAP = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ة": "ه", "ى": "ي", "ؤ": "و", "ئ": "ي"})
_ARABIC_MARKS = re.compile("[ً-ٰٟـ]")


def normalize(text):
    """تطبيع النص للبحث: حروف صغيرة، بدون تشكيل، وتوحيد الألف والتاء المربوطة."""
    return _ARABIC_MARKS.sub("", text.casefold()).translate(_ARABIC_MAP)


def natural_key(name):
    return [int(p) if p.isdigit() else p for p in re.split(r"(\d+)", normalize(name))]


def is_hidden(name):
    return name.startswith(".") or name.startswith("~$") or name.lower() in IGNORED_FILES


def list_entries(path):
    """يرجع (المجلدات, الملفات) مرتبة، بدون الملفات المخفية."""
    dirs, files = [], []
    try:
        with os.scandir(path) as it:
            for e in it:
                if is_hidden(e.name):
                    continue
                try:
                    (dirs if e.is_dir() else files).append(e.name)
                except OSError:
                    continue
    except OSError:
        pass
    dirs.sort(key=natural_key)
    files.sort(key=natural_key)
    return dirs, files


def subfolders(path):
    return list_entries(path)[0]


def validate_name(name):
    """يرجع رسالة خطأ بالعربية أو None إذا الاسم صالح."""
    if not name:
        return "الاسم فارغ."
    if INVALID_CHARS.search(name):
        return 'الاسم فيه رموز ممنوعة: < > : " / \\ | ? *'
    if name != name.rstrip(" ."):
        return "الاسم ما ينجمش ينتهي بنقطة أو فراغ."
    if name.split(".")[0].upper() in RESERVED_NAMES:
        return "هذا الاسم محجوز من النظام."
    return None


def search_names(root, query, limit=1000):
    """يبحث بالاسم (مع تجاهل الهمزات والتشكيل) في كل المجلدات.
    يرجّع (قائمة (مسار، مجلد؟)، النتائج مقطوعة؟)."""
    query = normalize(query.strip())
    results = []
    if not query or not root:
        return results, False
    for folder, dirs, files in os.walk(root):
        dirs[:] = sorted((d for d in dirs if not is_hidden(d)), key=natural_key)
        names = [(d, True) for d in dirs] + [(f, False) for f in sorted(files, key=natural_key) if not is_hidden(f)]
        for name, is_dir in names:
            if query in normalize(name):
                if len(results) >= limit:
                    return results, True
                results.append((os.path.join(folder, name), is_dir))
    return results, False


# ---------- تعمير النموذج (الكلمات المعلّمة) ----------
PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
_INVISIBLE = re.compile("[​-‏‪-‮⁦-⁩﻿]")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
# حروف ما يقبلهاش ملف Word (XML): حروف تحكّم مخفية تجي ساعات من ملفات PDF
_XML_INVALID = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ud800-\udfff\ufffe\uffff]")


def xml_safe(text):
    return _XML_INVALID.sub("", str(text)) if text is not None else ""


def placeholder_key(text):
    return normalize(_INVISIBLE.sub("", text).strip())


def fill_paragraph(p_el, mapping):
    """يبدّل {{...}} في فقرة. الكلمة المعلّمة ممكن تكون مقسومة على عدة runs:
    القيمة الجديدة تتكتب في أول run (فتبقى تنسيقاتو) والبقية تتفرّغ من الجزء المعلّم."""
    nodes = list(p_el.iter(qn("w:t")))
    originals = [t.text or "" for t in nodes]
    full = "".join(originals)
    if "{{" not in full:
        return
    matches = [m for m in PLACEHOLDER_RE.finditer(full) if placeholder_key(m.group(1)) in mapping]
    if not matches:
        return
    starts, pos = [], 0
    for text in originals:
        starts.append(pos)
        pos += len(text)
    texts = list(originals)
    for m in reversed(matches):  # من الآخر باش الإزاحات ما تتلخبطش
        start, end = m.span()
        value = mapping[placeholder_key(m.group(1))]
        first = True
        for i, original in enumerate(originals):
            a, b = starts[i], starts[i] + len(original)
            if not original or b <= start or a >= end:
                continue
            lo, hi = max(start, a) - a, min(end, b) - a
            texts[i] = texts[i][:lo] + (value if first else "") + texts[i][hi:]
            first = False
    for node, new, old in zip(nodes, texts, originals):
        if new != old:
            set_node_text(node, new)


def set_node_text(node, text):
    """يكتب النص في w:t، وكل سطر جديد (\\n) يتحوّل لفاصل سطر w:br داخل نفس الـ run."""
    first, *rest = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    node.text = first
    node.set(XML_SPACE, "preserve")
    prev = node
    for part in rest:
        br, t = OxmlElement("w:br"), OxmlElement("w:t")
        t.text = part
        t.set(XML_SPACE, "preserve")
        prev.addnext(br)
        br.addnext(t)
        prev = t


def fill_docx(src, dst, values):
    """ينسخ النموذج src إلى dst ويعمّر الكلمات المعلّمة في النص والجداول والـ header والـ footer."""
    if docx is None:
        raise RuntimeError("مكتبة python-docx غير مثبتة. نفّذ: pip install python-docx")
    values = {k: xml_safe(v) for k, v in values.items()}
    mapping = {placeholder_key(k): v for k, v in values.items()}
    document = docx.Document(src)
    present = set()
    for p_el in iter_paragraph_elements(document):  # الفقرات والجداول (حتى المتداخلة) وصناديق النص والـ header والـ footer
        text = "".join(t.text or "" for t in p_el.iter(qn("w:t")))
        present.update(placeholder_key(m.group(1)) for m in PLACEHOLDER_RE.finditer(text))
        fill_paragraph(p_el, mapping)
    # الأقسام اللي ما عندهاش {{...}} في النموذج: نعمّروها حسب عناوين خانات الجداول («الهدف المميز»، «رصد التصورات»...)
    filled = fill_by_labels(document, {k: v for k, v in values.items()
                                       if not section_in_template(k, present)})  # FillResult
    document.save(dst)
    return filled


# عناوين الخانات في نماذج المعلمين (بلا {{...}}) ← القسم اللي يتكتب بحذاها.
# قسم واحد ينجم يكون عندو برشا عناوين؛ وكي يلقى زوز عناوين لنفس القسم (رصد التصورات / بناء الفرضيات)
# الأول ياخو جزء الدليل والثاني ياخو جزء كتاب التلميذ.
TEMPLATE_LABELS = [
    (("المادة",), ["المادة"]),
    (("الدرس",), ["الدرس", "عنوان الدرس", "موضوع الدرس", "الموضوع"]),
    (("الأسبوع",), ["الأسبوع"]),
    (("التاريخ",), ["التاريخ"]),
    (("الوحدة",), ["الوحدة"]),
    (("الكفاية النهائية",), ["الكفاية النهائية", "نص الكفاية النهائية", "الكفاية"]),
    (("المكون الأول", "المكون الثاني"), ["مكون الكفاية", "مكونات الكفاية", "المكون"]),
    (("المكون الأول",), ["المكون الأول"]),
    (("المكون الثاني",), ["المكون الثاني"]),
    (("الهدف",), ["الهدف المميز", "الأهداف المميزة", "الهدف", "الأهداف", "الهدف المميز للحصة"]),
    (("المحتوى",), ["المحتوى", "المحتويات"]),
    (("المفاهيم",), ["المفاهيم"]),
    (("المستلزمات",), ["المستلزمات", "المستلزمات البيداغوجية"]),
    (("الحواجز",), ["الحواجز"]),
    (("مؤشرات التجاوز",), ["مؤشرات التجاوز"]),
    (("مؤشرات القدرة المستهدفة",), ["مؤشرات القدرة المستهدفة", "مؤشرات القدرة"]),
    (("المكتسبات السابقة",), ["تعهد المكتسبات", "المكتسبات السابقة", "أتعهد مكتسباتي", "المكتسبات"]),
    (("الوضعية المشكل",), ["الوضعية المشكل", "الوضعية المشكلة", "الوضعية الاستكشافية", "الوضعية الانطلاقية",
                           "ألاحظ وأتساءل", "الانطلاق"]),
    (("تحليل الوضعية ورصد التصورات",), ["تحليل الوضعية ورصد التصورات", "تحليل الوضعية", "رصد التصورات",
                                        "التصورات", "بناء الفرضيات", "الفرضيات", "أفترض"]),
    (("التحقق العلمي",), ["التحقق العلمي", "أجرب وأتثبت", "التجريب"]),
    (("الاستنتاج",), ["الاستنتاج", "أستنتج"]),
    (("التطبيق والتوظيف",), ["التطبيق والتوظيف", "التطبيق", "أطبق وأوظف", "التوظيف", "التعلم الآلي"]),
    (("التوسع والامتداد",), ["التوسع والامتداد", "التعلم الإدماجي", "الإدماج"]),
    (("التقييم",), ["التقييم", "التقويم", "أقيم تعلمي الجديد", "أقيم تعلمي"]),
]
# عنوان العمود اللي تتكتب فيه مراحل الحصة (كان الجدول فيه أعمدة: المراحل | نشاط المدرس | نشاط المتعلم...)
STAGE_HEADERS = ["المراحل", "المرحلة", "مراحل الدرس", "مراحل الحصة", "سير الدرس"]
ACTIVITY_HEADERS = ["نشاط المدرس", "نشاط المعلم", "نشاط المربي", "أنشطة المعلم", "سير الأنشطة", "سير الحصة", "الأنشطة"]
_BOOK_SPLIT = "\n\nكتاب التلميذ:\n"  # نفس الفاصل اللي يحطو build_generator_values


def _label_index():
    entries = [(_compact(alias)[0], keys) for keys, aliases in TEMPLATE_LABELS for alias in aliases]
    return sorted(entries, key=lambda e: -len(e[0]))


def _match_label(text, index):
    """خانة عنوان؟ يرجّع مفاتيح القسم. النص لازم يبدا بالعنوان كلمة كاملة («التحقق العلمي المنهج التجريبي...»)."""
    compact, _ = _compact(text.strip().rstrip(":：").strip())
    if not compact or len(compact) > 120:
        return None
    for key, keys in index:
        if compact == key or (compact.startswith(key) and not "\u0621" <= compact[len(key)] <= "\u064a"):
            return keys
        if compact.startswith(key) and len(key) >= 6 and len(compact) - len(key) > 8:
            return keys  # عنوان طويل متبوع بشرح («التحقق العلمي المنهج التجريبي / البحث الوثائقي»)
    return None


def _cell_text(tc):
    return "".join(t.text or "" for t in tc.iter(qn("w:t"))).strip()


def _grid_cells(tr):
    """خانات السطر مع موضع كل وحدة في شبكة الجدول (يحسب الخانات المدموجة أفقيا)."""
    cells, col = [], 0
    for tc in _own(tr, "w:tc", "w:tr"):
        span = tc.find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}")
        cells.append((col, tc))
        col += int(span.get(qn("w:val"))) if span is not None else 1
    return cells


def _write_cell(tc, text):
    """يكتب النص في الخانة (كل سطر في فقرة)، وكان الخانة فيها نص يزيدو تحتو."""
    paragraphs = tc.findall(qn("w:p"))
    empty = bool(paragraphs) and not _cell_text(tc)
    if empty:  # خانة فارغة (ساعات فيها برشا فقرات فارغة بعد الدمج): نكتبو من الفقرة الأولى
        for extra in paragraphs[1:]:
            tc.remove(extra)
        paragraphs = paragraphs[:1]
    base = paragraphs[-1] if paragraphs else None
    ppr = base.find(qn("w:pPr")) if base is not None else None
    for i, line in enumerate(text.replace("\r\n", "\n").split("\n")):
        if i == 0 and empty:
            p = base
        else:
            p = OxmlElement("w:p")
            if ppr is not None:
                p.append(copy.deepcopy(ppr))
            if paragraphs:
                paragraphs[-1].addnext(p)
            else:
                tc.append(p)
            paragraphs.append(p)
        p_ppr = p.find(qn("w:pPr"))
        if p_ppr is None:
            p_ppr = OxmlElement("w:pPr")
            p.insert(0, p_ppr)
        if p_ppr.find(qn("w:bidi")) is None:
            p_ppr.insert(0, OxmlElement("w:bidi"))  # فقرة عربية من اليمين لليسار
        run = OxmlElement("w:r")
        mark = p_ppr.find(qn("w:rPr"))  # تنسيق علامة الفقرة (الخط والحجم اللي اختارهم المعلم للخانة)
        rpr = copy.deepcopy(mark) if mark is not None else OxmlElement("w:rPr")
        if rpr.find(qn("w:rtl")) is None:
            rpr.append(OxmlElement("w:rtl"))
        run.append(rpr)
        t = OxmlElement("w:t")
        t.text = line
        t.set(XML_SPACE, "preserve")
        run.append(t)
        p.append(run)


class FillResult(set):
    """الأقسام اللي تكتبت في النموذج، و.empty = عناوين خانات لقيناها في النموذج أما القسم متاعها ما فيهش نص."""
    def __init__(self, filled=(), empty=()):
        super().__init__(filled)
        self.empty = list(empty)


def _own(el, tag, parent_tag):
    """العناصر tag اللي تابعين el مباشرة (حتى كان ملفوفين في w:sdt ولا w:customXml)، موش اللي في جدول داخلي."""
    out = []
    for child in el.iter(qn(tag)):
        up = child.getparent()
        while up is not None and up is not el and up.tag != qn(parent_tag):
            up = up.getparent()
        if up is el:
            out.append(child)
    return out


def fill_by_labels(document, values):
    """يعمّر نموذج ما فيهش {{...}}: كل خانة عنوان في جدول («الهدف المميز»، «رصد التصورات»...)
    يتكتب القسم متاعها في عمود «نشاط المدرس» كان موجود، وإلا في الخانة اللي بحذاها.
    يخدم في الجداول من اليمين لليسار ومن اليسار لليمين، وجدول يكمّل في الصفحة الجاية بلا سطر عناوين ياخو
    نفس أعمدة الجدول اللي قبلو. يرجّع FillResult."""
    index = _label_index()
    headers = {_compact(h)[0] for h in ACTIVITY_HEADERS}
    stage_headers = {_compact(h)[0] for h in STAGE_HEADERS}
    slots = []  # (مفاتيح القسم، الخانة اللي نكتبو فيها)
    layout = None  # (عدد الأعمدة، عمود المراحل، عمود نشاط المدرس) من آخر جدول فيه سطر عناوين
    for tbl in document.element.body.iter(qn("w:tbl")):
        rows = [_grid_cells(tr) for tr in _own(tbl, "w:tr", "w:tbl")]
        width = max((c[-1][0] + 1 for c in rows if c), default=0)
        current = layout if layout and layout[0] == width else None
        for cells in rows:
            texts = [_compact(_cell_text(tc))[0] for _, tc in cells]
            activity = next((col for (col, _), t in zip(cells, texts) if t in headers), None)
            if activity is not None:  # سطر العناوين: المراحل | نشاط المدرس | نشاط المتعلم | الوسائل
                stage = next((col for (col, _), t in zip(cells, texts) if t in stage_headers), None)
                current = layout = (width, stage, activity)
                continue
            positions = range(len(cells))
            if current and current[1] is not None:
                positions = [i for i, (col, _) in enumerate(cells) if col == current[1]]
            for pos in positions:
                keys = _match_label(_cell_text(cells[pos][1]), index)
                if not keys:
                    continue
                col = cells[pos][0]
                target = None
                if current and current[2] != col:
                    target = next((tc for gc, tc in cells if gc == current[2]), None)
                if target is None:  # بلا سطر عناوين: الخانة اللي بحذا العنوان (بعدو، وإلا قبلو كان العنوان في الآخر)
                    neighbor = pos + 1 if pos + 1 < len(cells) else pos - 1
                    target = cells[neighbor][1] if neighbor >= 0 else None
                if target is not None and not _match_label(_cell_text(target), index):
                    slots.append((keys, target))
                break  # خانة عنوان وحدة في السطر
    norm = {placeholder_key(k): (str(v).strip() if v is not None else "") for k, v in values.items()}
    by_section = {}
    for keys, target in slots:
        by_section.setdefault(keys, []).append(target)
    filled, empty = set(), []
    for keys, targets in by_section.items():
        if not any(placeholder_key(k) in norm for k in keys):
            continue  # القسم تعمّر بكلمة معلّمة {{...}} ولا ما يخصّش هالمذكرة
        parts = [norm.get(placeholder_key(k), "") for k in keys]
        text = "\n".join(p for p in parts if p)
        if not text:
            empty.append(" / ".join(keys))
            continue
        chunks = [text] * len(targets)  # كل خانة تتعمّر
        if len(targets) > 1 and _BOOK_SPLIT in text:  # عنوانين لنفس القسم: الدليل في الأول والكتاب في الثاني
            guide, book = (c.strip() for c in text.split(_BOOK_SPLIT, 1))
            chunks = [guide, book] + [text] * (len(targets) - 2)
        for target, chunk in zip(targets, chunks):
            if chunk:
                _write_cell(target, chunk)
        filled.update(k for k, p in zip(keys, parts) if p)
    return FillResult(filled, empty)


def auto_file_name(subject, lesson, week):
    parts = [subject.strip(), lesson.strip(), f"أسبوع {week}"]
    name = " - ".join(p for p in parts if p)
    name = re.sub(r"\s+", " ", INVALID_CHARS.sub(" ", name))
    return name.strip(" .")


# ---------- مولّد المذكرات (بلا إنترنت): استخراج النص من PDF وتقسيمه ----------
GUIDE_FILE, BOOK_FILE, KEYWORDS_FILE = "guide.pdf", "book.pdf", "keywords.json"
BOOK_KEY = "الكتاب"
STANDARD_KEYS = {"المادة", "الدرس", "الأسبوع", "التاريخ", BOOK_KEY}
# مفاتيح ملف keywords.json
SRC_GUIDE, SRC_BOOK = "الدليل", "الكتاب"
INLINE_KEY = "قيمة_في_نفس_السطر"  # عنوان: قيمة في نفس السطر (مثل «الوحدة : الهواء والتنفس»)
SEPARATORS_KEY = "عناوين_بدون_قسم"  # عناوين تقفل القسم الحالي بلا ما تبدا قسم جديد
SOURCE_LABELS = {"guide": "دليل المعلم", "book": "كتاب التلميذ"}

# مراحل الدرس حسب البرامج الرسمية التونسية: نفس المرحلة عندها اسم في دليل المعلم واسم في كتاب التلميذ.
# اسم القسم (المفتاح) هو الكلمة المعلّمة في النموذج: {{الوضعية المشكل}} ...
DEFAULT_KEYWORDS = {
    # --- بطاقة الجذاذة (دليل المعلم) ---
    "الكفاية النهائية": {SRC_GUIDE: ["نص الكفاية النهائية للمادة", "الكفاية النهائية للمادة", "الكفاية النهائية"],
                         SRC_BOOK: [], INLINE_KEY: True},
    "المكون الأول": {SRC_GUIDE: ["نص المكون الأول", "المكون الأول"], SRC_BOOK: [], INLINE_KEY: True},
    "المكون الثاني": {SRC_GUIDE: ["نص المكون الثاني", "المكون الثاني"], SRC_BOOK: [], INLINE_KEY: True},
    "الوحدة": {SRC_GUIDE: ["الوحدة"], SRC_BOOK: [], INLINE_KEY: True},
    "المفاهيم": {SRC_GUIDE: ["المفاهيم"], SRC_BOOK: ["المفاهيم"], INLINE_KEY: True},
    "المحتوى": {SRC_GUIDE: ["المحتوى"], SRC_BOOK: ["المحتوى"], INLINE_KEY: True},
    "الهدف": {SRC_GUIDE: ["الهدف المميز للوحدة", "الهدف المميز", "الهدف المميز للحصة", "الهدف"],
              SRC_BOOK: ["الهدف"], INLINE_KEY: True},
    "المستلزمات": {SRC_GUIDE: ["المستلزمات البيداغوجية", "المستلزمات", "الوسائل"], SRC_BOOK: [], INLINE_KEY: True},
    "الحواجز": {SRC_GUIDE: ["الحواجز"], SRC_BOOK: [], INLINE_KEY: True},
    "مؤشرات التجاوز": {SRC_GUIDE: ["مؤشرات التجاوز"], SRC_BOOK: []},
    "مؤشرات القدرة المستهدفة": {SRC_GUIDE: ["مؤشرات القدرة المستهدفة"], SRC_BOOK: [], INLINE_KEY: True},
    # --- مراحل الحصة ---
    "المكتسبات السابقة": {SRC_GUIDE: ["المكتسبات السابقة"], SRC_BOOK: ["أتعهد مكتسباتي السابقة", "أتعهد مكتسباتي"]},
    "الوضعية المشكل": {SRC_GUIDE: ["الوضعية المشكل", "الوضعية المشكلة", "الوضعية الانطلاقية"],
                       SRC_BOOK: ["ألاحظ وأتساءل"]},
    "تحليل الوضعية ورصد التصورات": {SRC_GUIDE: ["تحليل الوضعية ورصد التصورات", "تحليل الوضعية", "رصد التصورات"],
                                    SRC_BOOK: ["أفترض"]},
    # «النشاط الأول، النشاط الثاني...» تبقى ظاهرة داخل التحقق العلمي
    "التحقق العلمي": {SRC_GUIDE: ["التحقق العلمي", "النشاط"], SRC_BOOK: ["أجرب وأتثبت"]},
    "الاستنتاج": {SRC_GUIDE: ["الاستنتاج"], SRC_BOOK: ["أستنتج"]},
    "التطبيق والتوظيف": {SRC_GUIDE: ["التطبيق والتوظيف", "التطبيق"], SRC_BOOK: ["أطبق وأوظف"]},
    "التقييم": {SRC_GUIDE: ["التقييم", "التقويم"], SRC_BOOK: ["أقيم تعلمي الجديد", "أقيم تعلمي"]},
    "التوسع والامتداد": {SRC_GUIDE: ["التوسع والامتداد"], SRC_BOOK: []},
    "معجمي في العلوم": {SRC_GUIDE: [], SRC_BOOK: ["معجمي في العلوم"]},
    "أتهيأ لتعلمي اللاحق": {SRC_GUIDE: [], SRC_BOOK: ["أتهيأ لتعلمي اللاحق"]},
    SEPARATORS_KEY: ["التمشي البيداغوجي", "جذاذة تنشيط"],
}
# النسخة الأولى من keywords.json (قائمة كلمات لكل قسم): تتبدّل وحدها بالجديدة
_OLD_DEFAULT_KEYWORDS = {
    "الأهداف": ["الأهداف", "الأهداف التعلمية", "الكفاءة", "الكفاءات", "الهدف"],
    "الوضعية الانطلاقية": ["الوضعية الانطلاقية", "الوضعية المشكلة", "وضعية الانطلاق", "التمهيد"],
    "المراحل": ["المراحل", "مراحل الدرس", "سير الحصة", "سير الدرس", "سير الأنشطة"],
    "التقييم": ["التقييم", "التقويم", "تقييم", "تقويم"],
}
# كلمات عربية شائعة تُستعمل لكشف الحروف المعكوسة داخل الكلمة
COMMON_WORDS = {
    normalize(w) for w in (
        "في من على إلى الى عن أن ان التي الذي هذا هذه ذلك كل ثم أو او مع بين عند بعد قبل "
        "ما لا هو هي هم كان يكون تم حول خلال أي إذا اذا الدرس التلميذ المتعلم النشاط نشاط"
    ).split()
}
_ARABIC_CHAR = re.compile("[؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]")
_ARABIC_LETTERS = re.compile("[ء-ْ]+")
_CONTROL_CHARS = re.compile("[​-‏‪-‮⁦-⁩﻿]")
_LEADING_JUNK = re.compile(r"^[\W\d_]+")
_LIST_ITEM = re.compile(r"^(?:[-•·*–—▪●◦]|\d+\s*[-.)]|[ء-ي]\s*[-.)]|\(\s*\w{1,3}\s*\))\s*")
TEXT_MODES = [
    ("تلقائي", "auto"),
    ("ترتيب PDF الأصلي (بدون إعادة ترتيب)", "raw"),
    ("عكس حروف الكلمات", "letters"),
    ("عكس ترتيب الكلمات في السطر", "words"),
]


class MemoError(Exception):
    """خطأ بمعنى واضح للمستخدم (الرسالة بالعربية)."""


def load_keywords(subject_dir):
    """يقرأ keywords.json للمادة (ويكوّنو بالقيم الافتراضية إذا ما كانش موجود).
    يرجّع الإعدادات في شكل {"sections": {القسم: {"guide": [...], "book": [...], "inline": bool}}, "separators": [...]}."""
    path = os.path.join(subject_dir, KEYWORDS_FILE)
    data = None
    if os.path.isfile(path):
        try:
            with open(path, encoding="utf-8-sig") as f:
                data = json.load(f)
        except json.JSONDecodeError as e:
            raise MemoError(f"ملف {KEYWORDS_FILE} فيه غلطة في الكتابة (سطر {e.lineno}، عمود {e.colno}): {e.msg}\nصلّحو وعاود حاول.")
        except OSError as e:
            raise MemoError(f"ما نجمتش نقرا {KEYWORDS_FILE}:\n{e}")
    if data is None or data == _OLD_DEFAULT_KEYWORDS:
        data = DEFAULT_KEYWORDS
        try:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(DEFAULT_KEYWORDS, f, ensure_ascii=False, indent=2)
        except OSError as e:
            raise MemoError(f"ما نجمتش نكوّن ملف {KEYWORDS_FILE} في مجلد المادة:\n{e}")
    return validate_keywords(data)


def _word_list(value, where):
    if isinstance(value, str):
        value = [value]
    if not isinstance(value, list) or not all(isinstance(w, str) for w in value):
        raise MemoError(f"في {KEYWORDS_FILE}: {where} لازم تكون قائمة نصوص، مثلا [\"الوضعية المشكل\"].")
    return [w.strip() for w in value if w.strip()]


def validate_keywords(data):
    """يتثبّت من محتوى keywords.json. يقبل الشكل الجديد (الدليل/الكتاب) والشكل القديم (قائمة كلمات)."""
    if not isinstance(data, dict) or not data:
        raise MemoError(f"ملف {KEYWORDS_FILE} لازم يكون قاموس: اسم القسم ← الكلمات اللي تدل على عنوانو.")
    sections, separators = {}, []
    for name, spec in data.items():
        if name == SEPARATORS_KEY:
            separators = _word_list(spec, f"«{SEPARATORS_KEY}»")
            continue
        if name in STANDARD_KEYS:
            raise MemoError(f"في {KEYWORDS_FILE}: الاسم «{name}» محجوز، بدّلو (المحجوزة: {'، '.join(sorted(STANDARD_KEYS))}).")
        if isinstance(spec, dict):
            unknown = set(spec) - {SRC_GUIDE, SRC_BOOK, INLINE_KEY}
            if unknown:
                raise MemoError(
                    f"في {KEYWORDS_FILE}: القسم «{name}» فيه مفتاح غير معروف: {'، '.join(sorted(unknown))}.\n"
                    f"المفاتيح المقبولة: «{SRC_GUIDE}»، «{SRC_BOOK}»، «{INLINE_KEY}»."
                )
            guide = _word_list(spec.get(SRC_GUIDE, []), f"«{SRC_GUIDE}» في القسم «{name}»")
            book = _word_list(spec.get(SRC_BOOK, []), f"«{SRC_BOOK}» في القسم «{name}»")
            inline = spec.get(INLINE_KEY, False)
            if not isinstance(inline, bool):
                raise MemoError(f"في {KEYWORDS_FILE}: «{INLINE_KEY}» في القسم «{name}» لازم تكون true أو false.")
        else:  # الشكل القديم: نفس الكلمات للدليل والكتاب
            guide = book = _word_list(spec, f"قيمة القسم «{name}»")
            inline = False
        sections[name] = {"guide": guide, "book": book, "inline": inline}
    if not sections:
        raise MemoError(f"ملف {KEYWORDS_FILE} ما فيه حتى قسم.")
    return {"sections": sections, "separators": separators}


def keyword_vocab(config):
    """كلمات الأقسام + كلمات شائعة: تُستعمل لكشف النص العربي المعكوس."""
    vocab = set(COMMON_WORDS)
    for spec in config["sections"].values():
        for kw in spec["guide"] + spec["book"]:
            vocab.update(normalize(part) for part in kw.split())
    return vocab


def open_pdf(path):
    if fitz is None:
        raise MemoError("مكتبة PyMuPDF غير مثبتة. نفّذ: pip install PyMuPDF")
    if not os.path.isfile(path):
        raise MemoError(f"الملف غير موجود: {os.path.basename(path)}")
    try:
        doc = fitz.open(path)
    except Exception as e:  # noqa: BLE001
        raise MemoError(f"ما نجمتش نفتح الملف {os.path.basename(path)} (ربما تالف):\n{e}")
    if doc.needs_pass:
        doc.close()
        raise MemoError(f"الملف {os.path.basename(path)} محمي بكلمة سر. شيل الحماية وعاود حاول.")
    return doc


def pdf_page_count(path):
    doc = open_pdf(path)
    try:
        return doc.page_count
    finally:
        doc.close()


def _has_arabic(text):
    return bool(_ARABIC_CHAR.search(text))


_MIRROR = str.maketrans("()[]{}<>«»", ")(][}{><»«")
# أرقام وعبارات لاتينية داخل سطر عربي: تتكتب من اليسار لليمين
_LTR_RUN = re.compile(r"[0-9٠-٩]+(?:[.,:/][0-9٠-٩]+)*|[A-Za-z][A-Za-z0-9]*(?:[ .,:/'’&+-]+[A-Za-z0-9]+)*")


# ---------- خطوط AXt القديمة (Arabic XT متاع QuarkXPress) ----------
# برشة كتب مدرسية تونسية متعملة بـ QuarkXPress وخطوط AXt: كل شكل حرف عربي مخزّن في بلاصة حرف من
# ترميز Mac Roman (الألف بلاصة G...)، والسطر مخزّن بالترتيب المرئي (من اليسار لليمين).
# النص المستخرج يطلع رموز لاتينية (مثلا «á«Hô©dG» = «العربية»). الجدول هذا يرجّعها عربي.
# تبنى من دليل المعلم متاع الإيقاظ العلمي (المركز الوطني البيداغوجي) بمقارنة النص المشفّر بالنص الحقيقي.
AXT_TABLE = {
    "μ": "ك", "@": "",
    "'": "لا", "A": "ء", "B": "ٓ", "C": "ٔ", "D": "ٔ", "E": "ٕ", "F": "ئ", "G": "ا", "H": "ب", "I": "ة",
    "J": "ت", "K": "ث", "L": "ج", "M": "ح", "N": "خ", "O": "د", "P": "ذ", "Q": "ر", "R": "ز", "S": "س",
    "T": "ش", "U": "ص", "V": "ض", "W": "ط", "X": "ظ", "Y": "ع", "Z": "غ", "^": ",", "`": "يار", "a": "ف",
    "b": "ق", "c": "ك", "d": "ل", "e": "م", "f": "ن", "g": "ه", "h": "و", "i": "ى", "j": "ي", "o": "م",
    "z": "»", "{": "«", "¡": "ه", "£": "ط", "¤": "لى", "¥": "ق", "§": "ط", "¨": "غ", "©": "ع", "ª": "م",
    "«": "ي", "¬": "ه", "®": "ظ", "±": "ف", "´": "ع", "µ": "ك", "¶": "ظ", "º": "م", "»": "ي", "¿": "ن",
    "Á": "يم", "Ã": "بم", "Ä": "ئ", "Å": "ئ", "Æ": "غ", "È": "بر", "É": "ا", "Ê": "ني", "Ë": "يم", "Ì": "ثر",
    "Î": "تر", "Ñ": "ب", "Ò": "ير", "Ó": "لا", "Ö": "ب", "Ø": "ف", "Ù": "لمح", "Ú": "ين", "Û": "لمج",
    "Ü": "ب", "ß": "ظ", "à": "ت", "á": "ة", "â": "ت", "ã": "ث", "ä": "ت", "å": "ث", "æ": "ن", "ç": "ث",
    "è": "ج", "é": "ج", "ê": "ج", "ë": "ح", "ì": "ح", "í": "ح", "î": "خ", "ï": "خ", "ñ": "خ", "ò": "ذ",
    "ó": "د", "ô": "ر", "õ": "ز", "÷": "لج", "ø": "ن", "ù": "س", "û": "ش", "ü": "ص", "ÿ": "لخ", "ı": "لمخ",
    "Œ": "تج", "œ": "تج", "Ÿ": "لم", "ƒ": "و", "Ω": "م", "π": "ل", "–": "تح", "‘": "في", "’": "لا",
    "‚": "نج", "“": "تم", "”": "تم", "†": "ض", "‡": "مم", "•": "ط", "…": "ي", "‰": "نم", "‹": "لي",
    "›": "مج", "⁄": "لم", "™": "ع", "Ω": "م", "∂": "ك", "∏": "ل", "∑": "ك", "√": "ه", "∞": "ف", "∫": "ل",
    "≈": "ى", "≠": "غ", "≤": "ق", "≥": "ق", "◊": "لح", "ﬁ": "مح", "ﬂ": "مخ", "\"": "", "<": "", "\\": "",
    "n": "", "p": "", "q": "", "r": "", "s": "", "t": "", "u": "", "¢": "", "°": "", "·": "", "˘": "",
}
_AXT_LATIN_COMMON = set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789éèàùâêîôûçëïüÉÈÀÙ .,:;!?()-'\"")
_AXT_STRONG = set(AXT_TABLE) - _AXT_LATIN_COMMON  # رموز نادرة في النص الفرنسي ولا الإنقليزي
_HAMZA_MARKS = "\u0653\u0654\u0655"


def is_axt_text(text, font_name=""):
    """السطر مكتوب بخط AXt؟ (اسم الخط، وإلا نسبة الرموز الخاصة بالترميز)."""
    if "axt" in (font_name or "").lower():
        return True
    chars = [c for c in text if not c.isspace()]
    if len(chars) < 3 or _ARABIC_CHAR.search(text):
        return False
    strong = sum(1 for c in chars if c in _AXT_STRONG)
    return strong >= 2 and strong / len(chars) >= 0.2


def decode_axt(visual_text):
    """يحوّل سطر AXt (بالترتيب المرئي من اليسار لليمين) لنص عربي بالترتيب الصحيح."""
    text = visual_text.replace("fi", "ﬁ").replace("fl", "ﬂ")
    text = re.sub(r"(?<!\S)\\`(?!\S)", "•", text)  # «\`» وحدو = نقطة تعداد، ملصوق بكلمة = «يار»
    out = "".join(AXT_TABLE.get(ch, ch) for ch in reversed(text))
    out = out.translate(_MIRROR)
    out = _LTR_RUN.sub(lambda m: m.group(0)[::-1], out)  # الأرقام تقعد من اليسار لليمين
    # الهمزة والمدّة تتبع الألف/الواو/الياء اللي بعدها كي ما يكونش قبلها حرف يحملها
    out = re.sub(f"(?<![اويىأإآ])([{_HAMZA_MARKS}])([اويى])", r"\2\1", out)
    out = unicodedata.normalize("NFC", out)
    out = re.sub(r"(^|\s)ا (ل[اأإآ])", r"\1ا\2", out)  # «ا لأفقية» = «الأفقية»
    out = re.sub(r"(?<=[\u0621-\u064a]) (?=[ؤئ])", "", out)  # «م ؤشرات» = «مؤشرات»
    return re.sub(r"[ \t]{2,}", " ", out).strip()


def _spaced_ltr(row, h):
    """حروف السطر من اليسار لليمين مع الفراغات (حرف فراغ ولا مسافة كبيرة بين حرفين)."""
    text, prev = "", None
    for g in sorted(row, key=lambda g: g[1] + g[3]):
        if g[0].isspace():
            text += "" if text.endswith(" ") else " "
        else:
            if prev is not None and g[1] - prev[3] > 0.25 * h and not text.endswith(" "):
                text += " "
            text += g[0]
        prev = g
    return text


def page_lines(page, mode):
    """أسطر الصفحة بترتيب القراءة.
    نبنيو السطر حرف بحرف حسب بلاصة كل حرف في الصفحة (موش حسب ترتيب التخزين في PDF):
    هكا نتفادو الكلمات المعكوسة والمقطّعة، والأقواس والأرقام المقلوبة في النص العربي."""
    if mode == "raw":
        raw = page.get_text("text", sort=True).splitlines()
        lettered = [line for line in raw if any(ch.isalpha() for ch in line)]
        page_axt = bool(lettered) and sum(map(is_axt_text, lettered)) >= 0.5 * len(lettered)
        return [decode_axt(line) if is_axt_text(line) or (page_axt and not _ARABIC_CHAR.search(line)
                                                           and re.search("[A-Za-z]", line)) else line
                for line in raw]
    glyphs = []  # [نص، x0، y0، x1، y1، خط AXt؟]
    for block in page.get_text("rawdict")["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                axt_font = "axt" in span.get("font", "").lower()
                for c in span["chars"]:
                    ch, (x0, y0, x1, y1) = c["c"], c["bbox"]
                    if not axt_font and unicodedata.category(ch) in ("Mn", "Me") and glyphs:
                        glyphs[-1][0] += ch  # الشكل (الشدّة، الضمّة...) يتبع الحرف اللي قبلو
                    else:
                        glyphs.append([ch, x0, y0, x1, y1, axt_font])
    rows = []  # [مركز y، ارتفاع، حروف]
    for g in sorted(glyphs, key=lambda g: (g[2] + g[4]) / 2):
        yc, h = (g[2] + g[4]) / 2, max(g[4] - g[2], 1)
        if rows and abs(yc - rows[-1][0]) <= 0.5 * max(h, rows[-1][1]):
            rows[-1][2].append(g)
        else:
            rows.append([yc, h, [g]])
    # إذا أغلب أسطر الصفحة بخط AXt، الأسطر اللاتينية القصيرة (مثلا «IQGRh» = وزارة) زادة AXt
    lettered = [r for r in rows if any(ch.isalpha() for g in r[2] for ch in g[0])]
    axt_rows = sum(1 for r in lettered
                   if any(g[5] for g in r[2]) or is_axt_text("".join(g[0] for g in r[2])))
    page_axt = bool(lettered) and axt_rows >= 0.5 * len(lettered)
    lines = []
    for _yc, _h, row in rows:
        solid = [g for g in row if not g[0].isspace()]
        if not solid:
            continue
        heights = sorted(g[4] - g[2] for g in solid)
        h = max(heights[len(heights) // 2], 1)
        visual = "".join(g[0] for g in sorted(row, key=lambda g: g[1] + g[3]))
        if (any(g[5] for g in solid) or is_axt_text(visual)
                or (page_axt and not _ARABIC_CHAR.search(visual) and re.search("[A-Za-z]", visual))):
            lines.append(decode_axt(_spaced_ltr(row, h)))  # خط AXt: نفكّو الترميز
            continue
        arabic = sum(len(_ARABIC_CHAR.findall(g[0])) for g in solid)
        latin = sum(len(re.findall(r"[A-Za-z]", g[0])) for g in solid)
        rtl = arabic > latin
        ordered = sorted(row, key=lambda g: -(g[1] + g[3]) if rtl else (g[1] + g[3]))
        text, prev = "", None
        for g in ordered:
            if g[0].isspace():
                if text and not text.endswith(" "):
                    text += " "
                prev = g
                continue
            if prev is not None and text and not text.endswith(" "):
                gap = (prev[1] - g[3]) if rtl else (g[1] - prev[3])
                if gap > 0.25 * h:  # PDF بلا حرف فراغ: الفراغ يبان من المسافة
                    text += " "
            text += g[0]
            prev = g
        text = unicodedata.normalize("NFKC", text)
        if rtl:
            text = _LTR_RUN.sub(lambda m: m.group(0)[::-1], text.translate(_MIRROR))
        lines.append(text.strip())
    return lines


def looks_letter_reversed(lines, vocab):
    forward = backward = 0
    for line in lines:
        for token in re.findall(r"[ء-ْ]+", line):
            key = normalize(token)
            if key in vocab:
                forward += 3
            if key[::-1] in vocab:
                backward += 3
            if len(key) > 3 and key.startswith("ال"):
                forward += 1
            if len(key) > 3 and key.endswith("لا"):
                backward += 1
    return backward > forward * 1.5 and backward >= 3


# حروف من خطوط PDF مكسورة (رموز خاصة، حروف لغات أخرى بلاصة الشدّة والكسرة...): نحذفوها
_ODD_CHARS = re.compile(
    "[^\\s\u0000-\u024f\u0300-\u036f\u0370-\u03ff\u0600-\u06ff\u0750-\u077f\u08a0-\u08ff"
    "\u2000-\u2bff\ufb50-\ufdff\ufe70-\ufeff]"
)


def repair_lines(lines, mode, vocab):
    lines = [_ODD_CHARS.sub("", _CONTROL_CHARS.sub("", xml_safe(unicodedata.normalize("NFKC", line)))).strip() for line in lines]
    lines = [line for line in lines if line]
    if mode == "letters" or (mode == "auto" and looks_letter_reversed(lines, vocab)):
        lines = [_ARABIC_LETTERS.sub(lambda m: m.group(0)[::-1], line) for line in lines]
    if mode == "words":
        lines = [" ".join(reversed(line.split())) for line in lines]
    return lines


def extract_lines(pdf_path, first, last, mode, vocab):
    """أسطر النص للصفحات first..last (ترقيم من 1)."""
    doc = open_pdf(pdf_path)
    try:
        if not 1 <= first <= last <= doc.page_count:
            raise MemoError(f"مجال الصفحات غير صحيح: {first}-{last} (الملف {os.path.basename(pdf_path)} فيه {doc.page_count} صفحة).")
        raw = []
        for number in range(first - 1, last):
            raw.extend(page_lines(doc[number], mode))
    finally:
        doc.close()
    lines = repair_lines(raw, mode, vocab)
    if not lines:
        raise MemoError(
            f"ما لقيتش نص في الصفحات {first}-{last} من {os.path.basename(pdf_path)}.\n"
            "ربما الصفحات صور ممسوحة ضوئيا (Scan). التطبيق يقرا النص الحقيقي برك، وما يعملش OCR."
        )
    return lines


def count_headings(lines, config, source):
    """قدّاش من سطر عنوان (قسم أو فاصل) نلقاو في النص: مقياس لصحة القراية."""
    index = _keyword_index(config, source)
    return sum(1 for line in lines if match_heading(line, index, config))


MODE_LABELS = dict((value, label) for label, value in TEXT_MODES)


def extract_for_generator(pdf_path, first, last, mode, config, source):
    """يستخرج نص الصفحات. في الوضع «تلقائي» يجرّب كل طرق القراية ويختار اللي تلقى بيها أكثر عناوين
    (ملفات الكتب المدرسية ساعات النص فيها معكوس ولا مخزّن بترتيب مختلف).
    يرجّع (الأسطر، الطريقة المستعملة، عدد العناوين)."""
    vocab = keyword_vocab(config)
    modes = [mode] if mode != "auto" else ["auto", "raw", "letters", "words"]
    best = None
    for m in modes:
        lines = extract_lines(pdf_path, first, last, m, vocab)
        score = count_headings(lines, config, source)
        # طريقة أخرى غير العادية تربح كان إذا لقات عناوين أكثر بوضوح (موش بفارق عنوان صدفة)
        if best is None or (score > best[2] if best[1] != "auto" else score > best[2] * 1.5):
            best = (lines, m, score)
        if mode == "auto" and m == "auto" and score >= 3:
            break  # القراية العادية لقات عناوين: ما يلزمش نجرّبو غيرها
    return best


def extraction_notes(guide, book, guide_range, book_range):
    """ملاحظات للمعلم على القراية: الطريقة المختارة، وتحذير إذا ما تلقى حتى عنوان.
    guide/book: (الأسطر، الطريقة، عدد العناوين) أو None."""
    notes = []
    for label, info, pages in (("الدليل", guide, guide_range), ("الكتاب", book, book_range)):
        if not info:
            continue
        _lines, mode, found = info
        if mode != "auto":
            notes.append(f"{label}: النص تقرا بطريقة «{MODE_LABELS.get(mode, mode)}» (اختيار تلقائي، كانت الأحسن).")
        if found == 0:
            notes.append(
                f"⚠ ما لقيت حتى عنوان في {label} (الصفحات {pages[0]}-{pages[1]}). "
                "تثبّت من أرقام الصفحات (رقم الصفحة في ملف PDF، موش المكتوب على الورقة)، "
                f"وشوف «النص الأصلي: {label}» لوطة: إذا الحروف غريبة ولا مقطّعة، الـ PDF ما فيهوش نص عربي يتقرا."
            )
    return notes


def reflow(lines):
    """يلصق الأسطر اللي كانت مقطوعة بسبب عرض الصفحة، ويخلي القوائم والعناوين كل وحدة في سطر."""
    paragraphs, last_len = [], 0
    for line in lines:
        line = line.strip()
        if not line:  # سطر فارغ = فاصل بين فقرتين
            if paragraphs and paragraphs[-1]:
                paragraphs.append("")
            last_len = 0
            continue
        joinable = (
            paragraphs and paragraphs[-1] and last_len >= 45 and not _LIST_ITEM.match(line)
            and not re.search(r"[.:؟?!؛]$", paragraphs[-1])
        )
        if joinable:
            paragraphs[-1] += " " + line
        else:
            paragraphs.append(line)
        last_len = len(line)
    return "\n".join(paragraphs).strip()


def _compact(text):
    """نص مطبّع بلا فراغات + موضع كل حرف في النص الأصلي (باش «أجرب و أتثبت» = «أجرب وأتثبت»)."""
    chars, index = [], []
    for i, ch in enumerate(text):
        if ch.isspace():
            continue
        for nc in normalize(ch):
            chars.append(nc)
            index.append(i)
    return "".join(chars), index


_ORDINALS = "|".join(normalize(w) for w in (
    "الأول الثاني الثالث الرابع الخامس السادس السابع الثامن التاسع العاشر "
    "الحادي عشر الثاني عشر الأولى الثانية الثالثة الرابعة الخامسة السادسة"
).split())
_NUMBERED = re.compile(rf"^\W*(?:عدد\W*)?(?:\d+|[٠-٩]+|{_ORDINALS})\b")


def _keyword_index(config, source):
    entries = []
    for name, spec in config["sections"].items():
        for kw in spec[source]:
            key = _compact(kw)[0]
            if key:
                entries.append((key, name))
    for kw in config["separators"]:
        key = _compact(kw)[0]
        if key:
            entries.append((key, None))
    return sorted(entries, key=lambda e: -len(e[0]))  # الأطول أولا: «الهدف المميز» قبل «الهدف»


def match_heading(line, keyword_index, config):
    """إذا السطر عنوان يرجّع (القسم أو None للعناوين الفاصلة، المحتوى اللي في نفس السطر). وإلا None."""
    stripped = _LEADING_JUNK.sub("", line.strip())
    compact, index = _compact(stripped)
    for keyword, section in keyword_index:
        if not compact.startswith(keyword):
            continue
        end = index[len(keyword) - 1] + 1
        after = normalize(stripped[end:end + 1])
        if after and "\u0621" <= after[0] <= "\u064a":
            continue  # الكلمة المفتاحية لازم تكون كلمة كاملة: «المشكل» موش «المشكلة»
        rest = stripped[end:]
        inline = bool(section and config["sections"][section]["inline"])
        if _NUMBERED.match(normalize(rest)):
            # عنوان مرقّم (النشاط الأول، الاستنتاج 1، التطبيق (2)...): يبقى ظاهر داخل القسم
            return section, line.strip()
        colon = re.match(r"^[^:：]{0,40}[:：](.*)$", rest)
        if colon:
            content = colon.group(1)
        elif inline or (len(rest.strip()) <= 30 and not re.search(r"[.؟?!]", rest)):
            content = rest
        else:
            continue  # جملة عادية تبدا بنفس الكلمة، موش عنوان
        content = content.strip().lstrip("-–—:： \t")
        return section, content if re.search(r"[^\W\d_]", content) else ""
    return None


def segment_text(lines, config, source):
    """يقسّم أسطر الدليل (source="guide") أو الكتاب ("book") حسب العناوين.
    يرجّع (القسم ← نص، الأسطر اللي ما تبعت حتى قسم)."""
    index = _keyword_index(config, source)
    found = {name: [] for name in config["sections"]}
    current, unclassified = None, []
    for line in lines:
        heading = match_heading(line, index, config)
        if heading:
            current, content = heading
            if current is not None:
                if found[current]:
                    found[current].append("")  # كل ظهور جديد للقسم يبدا في فقرة جديدة
                if content:
                    found[current].append(content)
            continue
        (found[current] if current else unclassified).append(line)
    return {name: reflow(body) for name, body in found.items()}, unclassified


def build_generator_values(guide_lines, book_lines, config):
    """يحضّر قيم الخانات من الدليل والكتاب.
    يرجّع (خانات المراجعة القابلة للتعديل، قيم إضافية حسب المصدر: «القسم - الدليل» و«القسم - الكتاب»)."""
    guide, _ = segment_text(guide_lines, config, "guide")
    book, _ = segment_text(book_lines, config, "book") if book_lines else ({}, [])
    values, extras = {}, {}
    for name in config["sections"]:
        g, b = guide.get(name, ""), book.get(name, "")
        extras[f"{name} - {SRC_GUIDE}"] = g
        extras[f"{name} - {SRC_BOOK}"] = b
        if g and b:
            values[name] = f"{g}\n\n{SOURCE_LABELS['book']}:\n{b}"
        else:
            values[name] = g or b
    values[BOOK_KEY] = reflow(book_lines)
    return values, extras


def section_in_template(name, present):
    """القسم مستعمل في النموذج بأي شكل: {{القسم}} أو {{القسم - الدليل}} أو {{القسم - الكتاب}}."""
    return any(placeholder_key(k) in present for k in (name, f"{name} - {SRC_GUIDE}", f"{name} - {SRC_BOOK}"))


def iter_paragraph_elements(document):
    """كل فقرات المستند: الجسم (وفيه الجداول وصناديق النص) والـ header والـ footer."""
    roots = [document.element.body]
    for section in document.sections:
        for part in (
            section.header, section.first_page_header, section.even_page_header,
            section.footer, section.first_page_footer, section.even_page_footer,
        ):
            if not part.is_linked_to_previous:
                roots.append(part._element)
    for root in roots:
        yield from list(root.iter(qn("w:p")))


def template_placeholders(template_path):
    """أسماء الكلمات المعلّمة الموجودة في النموذج (مطبّعة)."""
    if docx is None:
        raise MemoError("مكتبة python-docx غير مثبتة. نفّذ: pip install python-docx")
    try:
        document = docx.Document(template_path)
    except Exception as e:  # noqa: BLE001
        raise MemoError(f"ما نجمتش نفتح النموذج {os.path.basename(template_path)} (ربما تالف):\n{e}")
    found = {}
    for p_el in iter_paragraph_elements(document):
        text = "".join(t.text or "" for t in p_el.iter(qn("w:t")))
        for m in PLACEHOLDER_RE.finditer(text):
            found[placeholder_key(m.group(1))] = m.group(1).strip()
    return found


def fmt_date(d):
    return f"{d.day():02d}/{d.month():02d}/{d.year()}"
