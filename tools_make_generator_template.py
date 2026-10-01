# -*- coding: utf-8 -*-
"""يكوّن generator_template.docx: نموذج جذاذة لمولّد المذكرات فيه كل الكلمات المعلّمة.
تشغيل: python tools_make_generator_template.py"""

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Traditional Arabic"
HEAD_FILL = "D9EAF7"
STAGE_FILL = "FDF1E6"


def rtl_paragraph(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p_pr = p._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.append(OxmlElement("w:bidi"))
    p.alignment = align
    return p


def run(p, text, bold=False, size=14, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = FONT
    r_pr = r._r.get_or_add_rPr()
    r_pr.get_or_add_rFonts().set(qn("w:cs"), FONT)
    r_pr.append(OxmlElement("w:rtl"))
    if bold:
        r_pr.append(OxmlElement("w:bCs"))
    sz = OxmlElement("w:szCs")
    sz.set(qn("w:val"), str(size * 2))
    r_pr.append(sz)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def rtl_table(table):
    tbl_pr = table._tbl.tblPr
    tbl_pr.append(OxmlElement("w:bidiVisual"))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"


def cell_text(cell, text, bold=False, size=13, align=WD_ALIGN_PARAGRAPH.RIGHT, color=None):
    p = cell.paragraphs[0]
    rtl_paragraph(p, align)
    run(p, text, bold=bold, size=size, color=color)


def main():
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.5))
    sec.top_margin = sec.bottom_margin = Cm(1.5)

    normal = d.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(13)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), FONT)

    # رأس وذيل الصفحة
    rtl_paragraph(sec.header.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
    run(sec.header.paragraphs[0], "{{المادة}} — {{الوحدة}}", size=11, color="555555")
    rtl_paragraph(sec.footer.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
    run(sec.footer.paragraphs[0], "{{الدرس}} — الأسبوع {{الأسبوع}} — {{التاريخ}}", size=10, color="777777")

    title = rtl_paragraph(d.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    run(title, "جذاذة درس: {{الدرس}}", bold=True, size=20, color="1F4E79")

    # بطاقة الجذاذة
    card = [
        ("المادة", "{{المادة}}"), ("الأسبوع / التاريخ", "الأسبوع {{الأسبوع}} — {{التاريخ}}"),
        ("الكفاية النهائية للمادة", "{{الكفاية النهائية}}"),
        ("المكوّن الأوّل", "{{المكون الأول}}"), ("المكوّن الثاني", "{{المكون الثاني}}"),
        ("الوحدة", "{{الوحدة}}"), ("الهدف المميّز", "{{الهدف}}"),
        ("المحتوى", "{{المحتوى}}"), ("المفاهيم", "{{المفاهيم}}"),
        ("المستلزمات البيداغوجيّة", "{{المستلزمات}}"), ("الحواجز", "{{الحواجز}}"),
        ("مؤشّرات التجاوز", "{{مؤشرات التجاوز}}"), ("مؤشّرات القدرة المستهدفة", "{{مؤشرات القدرة المستهدفة}}"),
    ]
    t = d.add_table(rows=len(card), cols=2)
    rtl_table(t)
    for i, (label, value) in enumerate(card):
        cell_text(t.cell(i, 0), label, bold=True)
        shade(t.cell(i, 0), HEAD_FILL)
        cell_text(t.cell(i, 1), value)
    for row in t.rows:
        row.cells[0].width, row.cells[1].width = Cm(5), Cm(13)

    rtl_paragraph(d.add_paragraph())
    h = rtl_paragraph(d.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    run(h, "سير الحصّة", bold=True, size=16, color="1F4E79")

    # المراحل: الاسم في الدليل / الاسم في كتاب التلميذ
    stages = [
        ("المكتسبات السابقة", "أتعهّد مكتسباتي السّابقة", "{{المكتسبات السابقة}}"),
        ("الوضعيّة المشكل", "ألاحظ وأتساءل", "{{الوضعية المشكل}}"),
        ("تحليل الوضعيّة ورصد التصوّرات", "أفترض", "{{تحليل الوضعية ورصد التصورات}}"),
        ("التحقّق العلمي", "أجرّب وأتثبّت", "{{التحقق العلمي}}"),
        ("الاستنتاج", "أستنتج", "{{الاستنتاج}}"),
        ("التطبيق والتوظيف", "أطبّق وأوظّف", "{{التطبيق والتوظيف}}"),
        ("التقييم", "أقيّم تعلّمي الجديد", "{{التقييم}}"),
        ("التوسّع والامتداد", "أتهيّأ لتعلّمي اللاّحق", "{{التوسع والامتداد}}\n{{أتهيأ لتعلمي اللاحق}}"),
    ]
    s = d.add_table(rows=len(stages) + 1, cols=2)
    rtl_table(s)
    cell_text(s.cell(0, 0), "المرحلة", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(s.cell(0, 1), "المحتوى والأنشطة", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade(s.cell(0, 0), HEAD_FILL)
    shade(s.cell(0, 1), HEAD_FILL)
    for i, (guide_name, book_name, value) in enumerate(stages, start=1):
        c = s.cell(i, 0)
        cell_text(c, guide_name, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        p = rtl_paragraph(c.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
        run(p, f"({book_name})", size=11, color="777777")
        shade(c, STAGE_FILL)
        cell_text(s.cell(i, 1), value)
    for row in s.rows:
        row.cells[0].width, row.cells[1].width = Cm(4.5), Cm(13.5)

    rtl_paragraph(d.add_paragraph())
    p = rtl_paragraph(d.add_paragraph())
    run(p, "معجمي في العلوم: ", bold=True)
    run(p, "{{معجمي في العلوم}}")

    d.save("generator_template.docx")
    print("generator_template.docx")


if __name__ == "__main__":
    main()
